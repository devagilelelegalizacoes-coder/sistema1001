"""Formulário de agendamento volante (DETRAN-RJ) preenchido a partir do lote.

O modelo em recursos/ é o mesmo Word que a equipe já usa — só a tabela de
veículos é reescrita a partir do banco. O resto do formulário (dados do
proprietário, observações, rodapé) fica como está no modelo.
"""
import copy
import io
from pathlib import Path

from docx import Document
from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from ..db import get_db
from ..models import Lote, Processo, Usuario
from ..seguranca import usuario_atual

router = APIRouter(prefix="/lotes", tags=["lotes"])

TEMPLATE = Path(__file__).resolve().parent.parent / "recursos" / "formulario_agendamento_volante.docx"


@router.get("")
def listar(db: Session = Depends(get_db), _: Usuario = Depends(usuario_atual)):
    lotes = db.scalars(select(Lote).order_by(Lote.criado_em.desc())).all()
    saida = []
    for lote in lotes:
        qtd = db.scalar(select(func.count(Processo.id)).where(Processo.lote_id == lote.id))
        saida.append({"id": lote.id, "nome": lote.nome, "tipo_servico": lote.tipo_servico,
                      "qtd_processos": qtd or 0})
    return saida


@router.get("/{lote_id}/formulario.docx")
def formulario(lote_id: int, db: Session = Depends(get_db), _: Usuario = Depends(usuario_atual)):
    lote = db.get(Lote, lote_id)
    if not lote:
        raise HTTPException(404, "Lote não encontrado")
    if not TEMPLATE.exists():
        raise HTTPException(500, "Modelo de formulário não está disponível no servidor")

    processos = db.scalars(
        select(Processo).where(Processo.lote_id == lote_id)
        .options(selectinload(Processo.veiculo))
        .order_by(Processo.id)
    ).all()
    if not processos:
        raise HTTPException(400, "Este lote não tem processos para preencher no formulário")

    doc = Document(TEMPLATE)
    tabela = doc.tables[0]
    linha_modelo = copy.deepcopy(tabela.rows[1]._tr)

    # remove todas as linhas de dados do modelo (mantém só o cabeçalho na linha 0)
    for linha in list(tabela.rows[1:]):
        linha._tr.getparent().remove(linha._tr)

    for p in processos:
        nova = copy.deepcopy(linha_modelo)
        tabela._tbl.append(nova)
        celulas = tabela.rows[-1].cells
        valores = [
            p.veiculo.placa or "",
            p.veiculo.renavam or "",
            p.duda_tp or "",
            p.data_venda.strftime("%d/%m/%Y") if p.data_venda else "",
            "",  # UF/origem — só usado em transferência de jurisdição, não modelado ainda
            p.numero_crv or "",
        ]
        for celula, valor in zip(celulas, valores):
            for paragrafo in celula.paragraphs:
                for run in paragrafo.runs[1:]:
                    run.text = ""
                if paragrafo.runs:
                    paragrafo.runs[0].text = valor
                else:
                    paragrafo.add_run(valor)

    buf = io.BytesIO()
    doc.save(buf)
    nome_arquivo = f"formulario_{lote.nome}.docx".replace(" ", "_")
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{nome_arquivo}"'},
    )
